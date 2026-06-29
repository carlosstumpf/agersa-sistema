"""
AGERSA - Banco Histórico
=========================
Importa cada Excel gerado pelo agersa_sistema.py
para um banco SQLite, acumulando histórico ao longo do tempo.

USO:
  python agersa_banco.py                    → importa Excel mais recente
  python agersa_banco.py --arquivo X.xlsx  → importa arquivo específico
  python agersa_banco.py --relatorio       → mostra resumo do banco
  python agersa_banco.py --exportar        → exporta histórico para Excel

BANCO: Downloads/AGERSA/agersa_historico.db
"""

import sqlite3
import pandas as pd
import openpyxl
import argparse
import sys
import glob
import os
from pathlib import Path
from datetime import datetime

# ═══════════════════════════════════════════════════════
#  CONFIGURAÇÕES
# ═══════════════════════════════════════════════════════
PASTA_AGERSA = Path.home() / "Downloads" / "AGERSA"
BANCO_PATH   = str(PASTA_AGERSA / "agersa_historico.db")

# ═══════════════════════════════════════════════════════
#  CONEXÃO E CRIAÇÃO DO BANCO
# ═══════════════════════════════════════════════════════
def conectar():
    PASTA_AGERSA.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(BANCO_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def criar_tabelas(conn):
    conn.executescript("""
    -- Registro de cada importação
    CREATE TABLE IF NOT EXISTS importacoes (
        id              INTEGER PRIMARY KEY AUTOINCREMENT,
        arquivo         TEXT NOT NULL,
        data_importacao TEXT NOT NULL,
        exercicio       INTEGER,
        mes_inicio      INTEGER,
        mes_fim         INTEGER,
        total_linhas    INTEGER,
        empenhado       REAL,
        liquidado       REAL,
        pago            REAL,
        a_pagar         REAL,
        rp_processado   REAL,
        rp_nao_proc     REAL,
        contratos       INTEGER,
        credores        INTEGER,
        notas_empenho   INTEGER
    );

    -- Dados de execução (granular por linha do CONSOLIDADO)
    CREATE TABLE IF NOT EXISTS fato_execucao (
        id                  INTEGER PRIMARY KEY AUTOINCREMENT,
        importacao_id       INTEGER REFERENCES importacoes(id),
        data_referencia     TEXT,
        exercicio           INTEGER,
        mes_referencia      INTEGER,

        -- Dimensões
        unidade_orcamentaria TEXT,
        unidade_gestora      TEXT,
        credor               TEXT,
        cnpj_credor          TEXT,
        nota_empenho         TEXT,
        historico            TEXT,
        processo_sei         TEXT,
        nota_liquidacao      TEXT,
        nf                   TEXT,
        programacao_desembolso TEXT,
        ordem_bancaria       TEXT,
        contrato_siafe       TEXT,
        contrato_original    TEXT,

        -- Valores de execução
        despesas_empenhadas  REAL DEFAULT 0,
        desp_emp_a_liquidar  REAL DEFAULT 0,
        despesas_liquidadas  REAL DEFAULT 0,
        desp_liq_a_pagar     REAL DEFAULT 0,
        despesas_a_pagar     REAL DEFAULT 0,
        despesas_pagas       REAL DEFAULT 0,

        -- Restos a pagar
        rp_nao_proc_a_liquidar    REAL DEFAULT 0,
        rp_nao_proc_inscritos     REAL DEFAULT 0,
        rp_nao_proc_ex_ant        REAL DEFAULT 0,
        rp_nao_proc_liq_a_pagar   REAL DEFAULT 0,
        rp_nao_proc_pagos         REAL DEFAULT 0,
        rp_proc_pagos             REAL DEFAULT 0,
        rp_proc_a_pagar           REAL DEFAULT 0,
        rp_proc_inscritos         REAL DEFAULT 0,
        rp_pagos                  REAL DEFAULT 0,
        retencoes_pagas           REAL DEFAULT 0,
        pd_paga_retencao          REAL DEFAULT 0
    );

    -- Índices para performance
    CREATE INDEX IF NOT EXISTS idx_fato_exercicio    ON fato_execucao(exercicio);
    CREATE INDEX IF NOT EXISTS idx_fato_importacao   ON fato_execucao(importacao_id);
    CREATE INDEX IF NOT EXISTS idx_fato_credor       ON fato_execucao(credor);
    CREATE INDEX IF NOT EXISTS idx_fato_contrato     ON fato_execucao(contrato_siafe);
    CREATE INDEX IF NOT EXISTS idx_fato_ne           ON fato_execucao(nota_empenho);

    -- View: resumo por importação
    CREATE VIEW IF NOT EXISTS vw_resumo_importacoes AS
    SELECT
        i.id,
        i.arquivo,
        i.data_importacao,
        i.exercicio,
        i.mes_inicio || ' a ' || i.mes_fim AS periodo,
        i.total_linhas,
        ROUND(i.empenhado, 2)     AS empenhado,
        ROUND(i.liquidado, 2)     AS liquidado,
        ROUND(i.pago, 2)          AS pago,
        ROUND(i.a_pagar, 2)       AS a_pagar,
        ROUND(i.pago / NULLIF(i.empenhado, 0) * 100, 1) AS pct_execucao
    FROM importacoes i
    ORDER BY i.data_importacao DESC;

    -- View: top credores
    CREATE VIEW IF NOT EXISTS vw_top_credores AS
    SELECT
        importacao_id,
        credor,
        ROUND(SUM(despesas_empenhadas), 2) AS empenhado,
        ROUND(SUM(despesas_liquidadas), 2) AS liquidado,
        ROUND(SUM(despesas_pagas), 2)      AS pago,
        COUNT(DISTINCT nota_empenho)        AS qtd_empenhos
    FROM fato_execucao
    GROUP BY importacao_id, credor
    ORDER BY empenhado DESC;

    -- View: evolução por contrato
    CREATE VIEW IF NOT EXISTS vw_contratos AS
    SELECT
        importacao_id,
        contrato_siafe,
        contrato_original,
        credor,
        ROUND(SUM(despesas_empenhadas), 2) AS empenhado,
        ROUND(SUM(despesas_liquidadas), 2) AS liquidado,
        ROUND(SUM(despesas_pagas), 2)      AS pago,
        ROUND(SUM(despesas_a_pagar), 2)    AS a_pagar
    FROM fato_execucao
    WHERE contrato_siafe IS NOT NULL AND contrato_siafe != ''
    GROUP BY importacao_id, contrato_siafe, contrato_original, credor
    ORDER BY empenhado DESC;
    """)
    conn.commit()
    print("✓ Banco criado/verificado.")

# ═══════════════════════════════════════════════════════
#  LEITURA DO EXCEL
# ═══════════════════════════════════════════════════════
def ler_excel(caminho):
    """Lê o Excel gerado pelo agersa_sistema.py e retorna dados estruturados."""
    print(f"Lendo: {Path(caminho).name}")
    wb = openpyxl.load_workbook(caminho, data_only=True)

    # ── Lê aba CONSOLIDADO ──────────────────────────────
    if 'CONSOLIDADO' not in wb.sheetnames:
        raise ValueError("Aba CONSOLIDADO não encontrada.")

    ws = wb['CONSOLIDADO']
    dados = []
    headers = []

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        # Linha 3 = cabeçalhos
        if i == 2:
            headers = [str(c).strip() if c else f"col_{i}" for i, c in enumerate(row)]
            continue
        # Linhas de dados (após cabeçalho, antes do TOTAL)
        if i > 2:
            if row[0] and str(row[0]).strip().upper() == 'TOTAL':
                break
            if any(v is not None for v in row):
                dados.append(dict(zip(headers, row)))

    df = pd.DataFrame(dados)
    print(f"  → {len(df)} linhas carregadas, {len(df.columns)} colunas")

    # ── Extrai metadados do subtítulo ──────────────────
    subtitulo = ws.cell(row=2, column=1).value or ""
    exercicio, mes_inicio, mes_fim = extrair_periodo(subtitulo, Path(caminho).name)

    # ── Lê totais do RESUMO ─────────────────────────────
    totais = ler_resumo(wb)

    return df, headers, exercicio, mes_inicio, mes_fim, totais

def extrair_periodo(subtitulo, nome_arquivo):
    """Extrai exercício e meses do subtítulo ou nome do arquivo."""
    import re
    exercicio = datetime.now().year
    mes_inicio = 1
    mes_fim = 12

    # Tenta pelo nome do arquivo: AGERSA_Execucao_2026_01a06_...
    m = re.search(r'(\d{4})_(\d{2})a(\d{2})', nome_arquivo)
    if m:
        exercicio  = int(m.group(1))
        mes_inicio = int(m.group(2))
        mes_fim    = int(m.group(3))
        return exercicio, mes_inicio, mes_fim

    # Tenta pelo subtítulo
    m = re.search(r'(\d{4})', subtitulo)
    if m:
        exercicio = int(m.group(1))

    return exercicio, mes_inicio, mes_fim

def ler_resumo(wb):
    """Extrai valores totais da aba RESUMO."""
    totais = {
        'empenhado': 0, 'liquidado': 0, 'pago': 0,
        'a_pagar': 0, 'rp_processado': 0, 'rp_nao_proc': 0,
        'contratos': 0, 'credores': 0, 'notas_empenho': 0,
    }
    if 'RESUMO' not in wb.sheetnames:
        return totais

    ws = wb['RESUMO']
    mapa = {
        'Despesas Empenhadas':             'empenhado',
        'Despesas Liquidadas':             'liquidado',
        'Despesas Pagas':                  'pago',
        'Despesas a Pagar':                'a_pagar',
        'Restos a Pagar Processados Inscritos': 'rp_processado',
        'Restos a Pagar Não Processados Inscritos': 'rp_nao_proc',
        'Contratos distintos':             'contratos',
        'Credores distintos':              'credores',
        'Notas de Empenho':                'notas_empenho',
    }
    for row in ws.iter_rows(values_only=True):
        if row[0] and str(row[0]).strip() in mapa:
            chave = mapa[str(row[0]).strip()]
            val = row[1]
            if isinstance(val, (int, float)):
                totais[chave] = val
    return totais

# ═══════════════════════════════════════════════════════
#  IMPORTAÇÃO
# ═══════════════════════════════════════════════════════
def importar(conn, caminho):
    """Importa um arquivo Excel para o banco."""

    # Verifica se já foi importado
    cur = conn.execute(
        "SELECT id FROM importacoes WHERE arquivo = ?",
        (Path(caminho).name,)
    )
    if cur.fetchone():
        print(f"  ⚠ Já importado anteriormente: {Path(caminho).name}")
        resp = input("  Reimportar? [s/N]: ").strip().lower()
        if resp != 's':
            return False
        conn.execute("DELETE FROM fato_execucao WHERE importacao_id = "
                     "(SELECT id FROM importacoes WHERE arquivo = ?)",
                     (Path(caminho).name,))
        conn.execute("DELETE FROM importacoes WHERE arquivo = ?",
                     (Path(caminho).name,))
        conn.commit()

    df, headers, exercicio, mes_inicio, mes_fim, totais = ler_excel(caminho)

    # Registra importação
    cur = conn.execute("""
        INSERT INTO importacoes
        (arquivo, data_importacao, exercicio, mes_inicio, mes_fim,
         total_linhas, empenhado, liquidado, pago, a_pagar,
         rp_processado, rp_nao_proc, contratos, credores, notas_empenho)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
    """, (
        Path(caminho).name,
        datetime.now().isoformat(),
        exercicio, mes_inicio, mes_fim,
        len(df),
        totais['empenhado'], totais['liquidado'],
        totais['pago'],      totais['a_pagar'],
        totais['rp_processado'], totais['rp_nao_proc'],
        totais['contratos'], totais['credores'], totais['notas_empenho'],
    ))
    importacao_id = cur.lastrowid

    # Mapeamento de colunas do Excel para campos do banco
    def col(nome):
        return df[nome] if nome in df.columns else pd.Series([None]*len(df))

    def num(serie):
        return pd.to_numeric(serie, errors='coerce').fillna(0)

    # Monta DataFrame para inserção
    fato = pd.DataFrame({
        'importacao_id':           importacao_id,
        'exercicio':               exercicio,
        'mes_referencia':          mes_inicio,
        'unidade_orcamentaria':    col('Unidade Orçamentária').astype(str),
        'unidade_gestora':         col('Unidade Gestora').astype(str),
        'credor':                  col('Credor').astype(str),
        'nota_empenho':            col('Nota de Empenho').astype(str),
        'historico':               col('Histórico').astype(str),
        'processo_sei':            col('Processo (NE)').astype(str),
        'nota_liquidacao':         col('Nota de Liquidação').astype(str),
        'nf':                      col('NF').astype(str),
        'programacao_desembolso':  col('Programação de Desembolso').astype(str),
        'ordem_bancaria':          col('Ordem Bancária').astype(str),
        'contrato_siafe':          col('Contrato SIAFE').astype(str),
        'contrato_original':       col('Contrato ORIGINAL').astype(str),
        'despesas_empenhadas':     num(col('Despesas Empenhadas')),
        'desp_emp_a_liquidar':     num(col('Desp. Emp. a Liquidar')),
        'despesas_liquidadas':     num(col('Despesas Liquidadas')),
        'desp_liq_a_pagar':        num(col('Desp. Liq. a Pagar')),
        'despesas_a_pagar':        num(col('Despesas a Pagar')),
        'despesas_pagas':          num(col('Despesas Pagas')),
        'rp_nao_proc_a_liquidar':  num(col('Restos a Pagar Não Processados a Liquidar')),
        'rp_nao_proc_inscritos':   num(col('Restos a Pagar Não Processados Inscritos')),
        'rp_nao_proc_ex_ant':      num(col('Restos a Pagar Não Processados Inscritos - Exercícios Anteriores')),
        'rp_nao_proc_liq_a_pagar': num(col('Restos a Pagar Não Processados Liquidados a Pagar')),
        'rp_nao_proc_pagos':       num(col('Restos a Pagar Não Processados Pagos')),
        'rp_proc_pagos':           num(col('Restos a Pagar Processados - Pagos')),
        'rp_proc_a_pagar':         num(col('Restos a Pagar Processados a Pagar')),
        'rp_proc_inscritos':       num(col('Restos a Pagar Processados Inscritos')),
        'rp_pagos':                num(col('RP Pagos')),
        'retencoes_pagas':         num(col('Retenções Pagas')),
        'pd_paga_retencao':        num(col('PD PAGA RETENÇÃO')),
    })

    # Substitui 'None' e 'nan' por None
    fato = fato.where(pd.notnull(fato), None)
    fato = fato.replace({'None': None, 'nan': None, 'NaT': None})

    # Insere em lote
    fato.to_sql('fato_execucao', conn, if_exists='append', index=False)
    conn.commit()

    print(f"  ✓ Importado: {len(df)} linhas → importacao_id={importacao_id}")
    print(f"    Exercício: {exercicio} | Período: {mes_inicio} a {mes_fim}")
    print(f"    Empenhado: R$ {totais['empenhado']:,.2f}")
    print(f"    Pago:      R$ {totais['pago']:,.2f}")
    return True

# ═══════════════════════════════════════════════════════
#  RELATÓRIOS DO BANCO
# ═══════════════════════════════════════════════════════
def mostrar_relatorio(conn):
    print("\n" + "="*60)
    print("  BANCO HISTÓRICO AGERSA")
    print("="*60)

    # Importações
    rows = conn.execute("SELECT * FROM vw_resumo_importacoes").fetchall()
    if not rows:
        print("  Nenhuma importação encontrada.")
        return

    print(f"\n{'ID':<4} {'Arquivo':<35} {'Período':<12} {'Empenhado':>16} {'Pago':>16} {'%':>6}")
    print("-"*95)
    for r in rows:
        print(f"{r['id']:<4} {r['arquivo'][:34]:<35} {r['periodo']:<12} "
              f"R$ {r['empenhado']:>13,.2f} R$ {r['pago']:>13,.2f} {r['pct_execucao']:>5.1f}%")

    # Top 10 credores da última importação
    ultima_id = rows[0]['id']
    print(f"\n  TOP 10 CREDORES (importação {ultima_id}):")
    print(f"  {'Credor':<50} {'Empenhado':>16} {'Pago':>16}")
    print("  " + "-"*84)
    credores = conn.execute(
        "SELECT credor, empenhado, pago FROM vw_top_credores "
        "WHERE importacao_id=? LIMIT 10", (ultima_id,)
    ).fetchall()
    for c in credores:
        nome = str(c['credor'])[:49]
        print(f"  {nome:<50} R$ {c['empenhado']:>13,.2f} R$ {c['pago']:>13,.2f}")

    print(f"\n  Banco: {BANCO_PATH}")
    print("="*60)

def exportar_historico(conn):
    """Exporta histórico completo para Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    nome = f"AGERSA_Historico_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    caminho = str(PASTA_AGERSA / nome)

    wb = Workbook()

    # Aba 1: Importações
    ws1 = wb.active
    ws1.title = "Importações"
    rows = conn.execute("SELECT * FROM vw_resumo_importacoes").fetchall()
    if rows:
        ws1.append(list(rows[0].keys()))
        for r in rows:
            ws1.append(list(r))

    # Aba 2: Top Credores
    ws2 = wb.create_sheet("Top Credores")
    rows2 = conn.execute("SELECT * FROM vw_top_credores LIMIT 500").fetchall()
    if rows2:
        ws2.append(list(rows2[0].keys()))
        for r in rows2:
            ws2.append(list(r))

    # Aba 3: Contratos
    ws3 = wb.create_sheet("Contratos")
    rows3 = conn.execute("SELECT * FROM vw_contratos LIMIT 500").fetchall()
    if rows3:
        ws3.append(list(rows3[0].keys()))
        for r in rows3:
            ws3.append(list(r))

    # Aba 4: Evolução mensal
    ws4 = wb.create_sheet("Evolução Mensal")
    rows4 = conn.execute("""
        SELECT
            exercicio, mes_referencia,
            ROUND(SUM(despesas_empenhadas),2) AS empenhado,
            ROUND(SUM(despesas_liquidadas),2) AS liquidado,
            ROUND(SUM(despesas_pagas),2)      AS pago,
            ROUND(SUM(despesas_a_pagar),2)    AS a_pagar,
            COUNT(DISTINCT nota_empenho)       AS qtd_ne,
            COUNT(DISTINCT credor)             AS qtd_credores
        FROM fato_execucao
        GROUP BY exercicio, mes_referencia
        ORDER BY exercicio, mes_referencia
    """).fetchall()
    if rows4:
        ws4.append(list(rows4[0].keys()))
        for r in rows4:
            ws4.append(list(r))

    wb.save(caminho)
    print(f"✓ Histórico exportado: {caminho}")
    return caminho

# ═══════════════════════════════════════════════════════
#  INTEGRAÇÃO COM O PIPELINE
# ═══════════════════════════════════════════════════════
def importar_mais_recente():
    """Importa o Excel mais recente da pasta AGERSA."""
    arquivos = sorted(
        PASTA_AGERSA.glob("AGERSA_Execucao_*.xlsx"),
        key=os.path.getmtime,
        reverse=True
    )
    if not arquivos:
        print(f"Nenhum arquivo encontrado em: {PASTA_AGERSA}")
        return False

    conn = conectar()
    criar_tabelas(conn)
    resultado = importar(conn, str(arquivos[0]))
    conn.close()
    return resultado

# ═══════════════════════════════════════════════════════
#  MAIN
# ═══════════════════════════════════════════════════════
def main():
    parser = argparse.ArgumentParser(description='AGERSA - Banco Histórico')
    parser.add_argument('--arquivo',   help='Importar arquivo Excel específico')
    parser.add_argument('--relatorio', action='store_true', help='Mostrar relatório do banco')
    parser.add_argument('--exportar',  action='store_true', help='Exportar histórico para Excel')
    parser.add_argument('--todos',     action='store_true', help='Importar todos os arquivos da pasta')
    args = parser.parse_args()

    conn = conectar()
    criar_tabelas(conn)

    if args.relatorio:
        mostrar_relatorio(conn)

    elif args.exportar:
        exportar_historico(conn)

    elif args.todos:
        arquivos = sorted(PASTA_AGERSA.glob("AGERSA_Execucao_*.xlsx"),
                          key=os.path.getmtime)
        print(f"Importando {len(arquivos)} arquivo(s)...")
        for arq in arquivos:
            importar(conn, str(arq))

    elif args.arquivo:
        importar(conn, args.arquivo)

    else:
        # Sem argumentos: importa o mais recente
        importar_mais_recente()
        mostrar_relatorio(conn)

    conn.close()

if __name__ == "__main__":
    main()
